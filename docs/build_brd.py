"""
Generates the BridgeIQ Business Requirements Document (BRD) as a Word .docx file
Run: python build_brd.py
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime
import os

OUT = os.path.join(os.path.dirname(__file__), "BridgeIQ_BRD.docx")

NAVY  = RGBColor(0x1A, 0x3A, 0x5C)
BLUE  = RGBColor(0x4F, 0x8E, 0xF7)
GRAY  = RGBColor(0x5D, 0x6D, 0x7E)
BLACK = RGBColor(0x0D, 0x1B, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# ── Page margins ─────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

# ── Helpers ───────────────────────────────────────────────────────────────────
def heading1(text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '1A3A5C')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def heading2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = BLUE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    return p

def body(text, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size  = Pt(11)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(4)
    return p

def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(11)
    p.paragraph_format.left_indent = Inches(0.3 * (level + 1))
    p.paragraph_format.space_after = Pt(2)
    return p

def table_row(table, cells, bold=False, header=False):
    row = table.add_row()
    for i, cell_text in enumerate(cells):
        cell = row.cells[i]
        cell.text = str(cell_text)
        for para in cell.paragraphs:
            para.runs[0].font.size = Pt(10)
            para.runs[0].font.bold = bold or header
            if header:
                para.runs[0].font.color.rgb = WHITE
        if header:
            cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), '1A3A5C')
            cell._tc.get_or_add_tcPr().append(shd)
    return row

# ── Cover Page ────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("\n\n\nBRIDGEIQ")
run.font.size  = Pt(36)
run.font.bold  = True
run.font.color.rgb = NAVY

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Business Requirements Document")
r2.font.size  = Pt(18)
r2.font.color.rgb = BLUE

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run("AI-Powered Customer Intelligence & Process Transformation\nApex Solutions B2B SaaS Platform")
r3.font.size  = Pt(13)
r3.font.color.rgb = GRAY
r3.font.italic = True

doc.add_paragraph("\n\n")

meta_table = doc.add_table(rows=1, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
rows_data = [
    ("Document Version", "1.0"),
    ("Status", "Final"),
    ("Date", datetime.date.today().strftime("%B %d, %Y")),
    ("Author", "Sai Hemanth — Technical Business Analyst"),
    ("Reviewed By", "Product Owner, Engineering Lead, CS Director"),
    ("Confidentiality", "Internal Use Only"),
]
for label, value in rows_data:
    row = meta_table.add_row()
    row.cells[0].text = label
    row.cells[1].text = value
    for cell in row.cells:
        for para in cell.paragraphs:
            if para.runs:
                para.runs[0].font.size = Pt(10.5)
    row.cells[0].paragraphs[0].runs[0].font.bold = True

doc.add_page_break()

# ── 1. Executive Summary ──────────────────────────────────────────────────────
heading1("1. Executive Summary")
body(
    "Apex Solutions, a B2B SaaS company delivering project management software to mid-size enterprises, "
    "is experiencing compounding operational challenges that directly threaten revenue retention and growth. "
    "An internal analysis of operational data revealed three critical failure areas: a 22% customer churn rate "
    "driven by poor onboarding completion and low product engagement, a support backlog with SLA breach rates "
    "exceeding 30% for high-priority tickets, and no automated early warning system to identify at-risk customers "
    "before they churn."
)
body(
    "This Business Requirements Document defines the scope, objectives, functional requirements, and success "
    "criteria for BridgeIQ — an AI-powered customer intelligence and process transformation platform. "
    "BridgeIQ integrates structured data analytics, AI-driven customer feedback analysis, and automated "
    "early warning capabilities to reduce churn by 35%, improve onboarding completion rates from 82% to 95%, "
    "and reduce average support resolution time by 40% within 12 months of deployment."
)

# ── 2. Business Context ───────────────────────────────────────────────────────
heading1("2. Business Context & Problem Statement")

heading2("2.1 Company Background")
body(
    "Apex Solutions was founded in 2018 and serves 500+ B2B customers across 10 industries including "
    "Healthcare, Finance, Retail, and Manufacturing. The company operates on a subscription model with "
    "three plan tiers — Starter ($299-499/mo), Growth ($999-2,499/mo), and Enterprise ($4,999-12,999/mo) — "
    "generating a combined Monthly Recurring Revenue (MRR) exceeding $1.2M."
)

heading2("2.2 Problem Statement")
body("Data analysis has identified the following quantified business problems:")
problems = [
    ("Customer Churn", "22% annual churn rate — industry benchmark is 5-7% for B2B SaaS. Churned customers have a measurably lower health score (avg 58.3 vs 82.1 for active) but no automated signal exists to flag them before they cancel."),
    ("Onboarding Failure", "18% of customers do not complete onboarding. Incomplete onboarding correlates with a 2.4x higher churn rate. Primary blockers: IT approval delays, data quality issues, and stakeholder unavailability — all manageable with a structured process."),
    ("Support Overload", "4,000+ support tickets analyzed show an average resolution time of 28 hours. SLA breach rate for Critical tickets exceeds 35%. Billing and Technical Issue categories account for 48% of all tickets."),
    ("No Predictive Intelligence", "Customer health scores are computed manually. No automated system exists to aggregate churn signals (low usage, high ticket volume, poor CSAT, incomplete onboarding) into a composite risk score."),
    ("Unstructured Feedback", "Customer feedback, survey responses, and ticket descriptions are unread at scale. High-value product insights and process failures are buried in unstructured text with no analysis pipeline."),
]
for title, desc in problems:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{title}: ").font.bold = True
    p.add_run(desc).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)

# ── 3. Project Scope ──────────────────────────────────────────────────────────
heading1("3. Project Scope")

heading2("3.1 In Scope")
in_scope = [
    "AI-powered customer health scoring engine with composite risk signal (usage, tickets, onboarding, NPS)",
    "Automated at-risk customer alerting via Slack / email for Customer Success team",
    "Structured onboarding workflow redesign with stage-gated milestones and blocker tracking",
    "Executive analytics dashboard (MRR, churn, onboarding completion, support KPIs)",
    "AI feedback analysis pipeline — NLP-based categorization of support tickets and survey responses",
    "AI requirements generation tool for internal BA / product team use",
    "SQL analytics layer — 10+ standardized queries for business reporting",
    "UAT framework and user acceptance test plan",
]
for item in in_scope:
    bullet(item)

heading2("3.2 Out of Scope")
out_scope = [
    "CRM system replacement or migration (integration only)",
    "Native mobile application",
    "Third-party billing system changes",
    "Custom enterprise SSO implementation (Phase 2)",
    "Real-time data streaming infrastructure (Phase 2)",
]
for item in out_scope:
    bullet(item)

# ── 4. Stakeholders ───────────────────────────────────────────────────────────
heading1("4. Stakeholder Analysis")
body("The following stakeholders have been identified through requirements gathering sessions and impact assessment:")

sh_table = doc.add_table(rows=1, cols=5)
sh_table.style = "Table Grid"
headers = ["Stakeholder", "Role", "Interest", "Influence", "Engagement Strategy"]
table_row(sh_table, headers, header=True)

stakeholders = [
    ("VP Customer Success", "Project Sponsor", "Churn reduction, CS team efficiency", "High", "Weekly status, co-design sessions"),
    ("Product Manager", "Key Decision Maker", "Feature roadmap alignment, user story prioritization", "High", "Sprint planning, backlog grooming"),
    ("Engineering Lead", "Technical Authority", "System architecture, API integration feasibility", "High", "Technical review sessions"),
    ("CS Team (6 CSMs)", "Primary End User", "Usable dashboards, actionable alerts", "Medium", "User interviews, UAT testing"),
    ("CFO", "Financial Approver", "ROI, MRR impact, cost justification", "High", "Executive summary, milestone reviews"),
    ("Support Team Lead", "Key User", "Ticket workflow, SLA visibility", "Medium", "Process mapping workshops"),
    ("Data Engineering", "Technical Implementer", "Data pipeline, SQL schema, data quality", "Medium", "Design sessions, code reviews"),
]
for row_data in stakeholders:
    table_row(sh_table, row_data)

doc.add_paragraph()

# ── 5. Functional Requirements ────────────────────────────────────────────────
heading1("5. Functional Requirements")

heading2("FR-01: Customer Health Scoring Engine")
reqs_01 = [
    "FR-01.1: System shall compute a composite health score (0-100) for each active customer on a daily basis.",
    "FR-01.2: Health score shall incorporate: product usage frequency, feature breadth, support ticket volume, CSAT score, onboarding completion status, and NPS score.",
    "FR-01.3: Customers with a composite risk score ≥50 shall be flagged as 'At Risk'.",
    "FR-01.4: Risk score thresholds shall be configurable by CS Director without code changes.",
]
for r in reqs_01:
    bullet(r)

heading2("FR-02: Automated Early Warning Alerts")
reqs_02 = [
    "FR-02.1: System shall send a Slack notification to the assigned CSM when a customer's risk score increases by ≥15 points within 7 days.",
    "FR-02.2: Alert shall include: customer name, plan, MRR, risk score, top 3 contributing risk factors, and a direct link to the customer dashboard.",
    "FR-02.3: Daily digest email summarizing all at-risk customers shall be sent to CS Manager at 8:00 AM EST.",
    "FR-02.4: Alerts shall not duplicate — a customer triggering multiple conditions within 24 hours shall generate one consolidated alert.",
]
for r in reqs_02:
    bullet(r)

heading2("FR-03: Onboarding Workflow Redesign")
reqs_03 = [
    "FR-03.1: System shall provide a digital onboarding checklist with 5 structured stages: Account Setup, Data Migration, User Training, Integration Config, Go-Live.",
    "FR-03.2: Each stage shall have defined completion criteria, assigned owner, and SLA (target completion days).",
    "FR-03.3: Customers stuck in a stage for more than 5 business days shall trigger an automatic escalation to their CSM.",
    "FR-03.4: CS team shall be able to log blockers (drop-down + free text) per stage with timestamped audit trail.",
]
for r in reqs_03:
    bullet(r)

heading2("FR-04: Executive Analytics Dashboard")
reqs_04 = [
    "FR-04.1: Dashboard shall display: Active Customers, Total MRR, ARR, Churn Rate, Average Health Score, Open Tickets, Onboarding Completion Rate.",
    "FR-04.2: All KPI cards shall show delta vs. prior period (MoM or QoQ selectable).",
    "FR-04.3: Revenue trend chart shall display monthly net revenue for trailing 24 months.",
    "FR-04.4: Dashboard shall be filterable by: plan type, region, industry, and date range.",
    "FR-04.5: At-risk customer table shall be sortable by risk score, MRR, and last activity date.",
]
for r in reqs_04:
    bullet(r)

heading2("FR-05: AI Feedback Analysis Pipeline")
reqs_05 = [
    "FR-05.1: System shall accept unstructured text input (support tickets, survey responses, feedback forms) and return structured analysis.",
    "FR-05.2: Output shall include: sentiment classification, pain point ranking by business impact, root cause categories, and auto-generated user stories.",
    "FR-05.3: User stories generated by AI shall follow Agile format: As a [persona], I want [feature], so that [outcome].",
    "FR-05.4: Users shall be able to download the analysis report as a .txt or .docx file.",
]
for r in reqs_05:
    bullet(r)

# ── 6. Non-Functional Requirements ───────────────────────────────────────────
heading1("6. Non-Functional Requirements")

nfr_table = doc.add_table(rows=1, cols=4)
nfr_table.style = "Table Grid"
table_row(nfr_table, ["Category", "Requirement", "Target", "Priority"], header=True)

nfrs = [
    ("Performance", "Dashboard page load time", "< 3 seconds for any page", "High"),
    ("Performance", "AI analysis response time", "< 15 seconds per request", "High"),
    ("Availability", "Platform uptime SLA", "99.5% monthly uptime", "High"),
    ("Security", "Data encryption", "All data encrypted in transit (TLS 1.3) and at rest (AES-256)", "Critical"),
    ("Security", "Authentication", "SSO / OAuth 2.0 support with MFA enforcement", "High"),
    ("Scalability", "Customer data volume", "Support up to 5,000 customers without performance degradation", "Medium"),
    ("Compliance", "Data residency", "Customer data processed and stored within US/EU regions per contract", "High"),
    ("Usability", "Onboarding new users", "New CS team member productive within 2 hours, no training required", "Medium"),
    ("Maintainability", "SQL query updates", "New KPI reports deployable without engineering involvement", "Medium"),
]
for row_data in nfrs:
    table_row(nfr_table, row_data)

doc.add_paragraph()

# ── 7. Process Design ─────────────────────────────────────────────────────────
heading1("7. Process Design — AS-IS vs TO-BE")

heading2("7.1 AS-IS: Customer Churn Detection Process")
as_is = [
    "CSM manually reviews customer accounts weekly (4+ hours/CSM/week)",
    "Health data pulled from 3 separate spreadsheets — no single source of truth",
    "No standardized risk criteria — entirely dependent on individual CSM judgment",
    "At-risk customers often identified only after cancellation notice received",
    "Escalation path informal — email chains with no SLA or audit trail",
]
for item in as_is:
    bullet(item)

body("Key Pain Points: Manual effort, inconsistent criteria, reactive (not proactive), no visibility for management.", italic=True, color=GRAY)

heading2("7.2 TO-BE: Automated Customer Intelligence Process")
to_be = [
    "Daily automated health score computation from 6 data signals — zero manual effort",
    "Unified dashboard — single source of truth accessible to all CS stakeholders",
    "Rule-based + AI composite risk scoring with configurable thresholds",
    "Proactive Slack/email alert sent to CSM when risk score crosses threshold",
    "Structured escalation workflow with SLA tracking and audit log",
    "Weekly AI-generated insight report from customer feedback batch processed overnight",
]
for item in to_be:
    bullet(item)

body("Expected Outcomes: 35% churn reduction, 4 hours/week saved per CSM, proactive intervention at scale.", italic=True, color=BLUE)

# ── 8. Assumptions & Constraints ─────────────────────────────────────────────
heading1("8. Assumptions & Constraints")

heading2("8.1 Assumptions")
assumptions = [
    "Customer data is available in a centralized SQL-accessible database at project kick-off.",
    "Slack is the primary communication tool for the CS team and integration access will be granted.",
    "Engineering team has capacity to allocate 2 developers for the 3-month build phase.",
    "All stakeholders are available for bi-weekly review sessions throughout the project.",
    "AI API (Anthropic Claude) access and billing is approved by the CFO prior to development start.",
]
for a in assumptions:
    bullet(a)

heading2("8.2 Constraints")
constraints = [
    "Budget cap of $120,000 for Phase 1 (build + deploy). Phase 2 features require separate approval.",
    "All deliverables must be production-ready within 12 weeks of project kick-off.",
    "No changes to the existing billing system infrastructure during Phase 1.",
    "Customer PII must remain within the existing data governance framework — no third-party data sharing.",
]
for c in constraints:
    bullet(c)

# ── 9. Success Metrics ────────────────────────────────────────────────────────
heading1("9. Success Metrics & KPIs")

sm_table = doc.add_table(rows=1, cols=4)
sm_table.style = "Table Grid"
table_row(sm_table, ["Metric", "Current Baseline", "Target (12 months)", "Measurement Method"], header=True)

metrics = [
    ("Annual Churn Rate", "22%", "≤ 14%", "CRM data — monthly cohort analysis"),
    ("Onboarding Completion Rate", "82%", "≥ 95%", "Onboarding tracker — stage completion data"),
    ("Avg Support Resolution Time", "28 hours", "≤ 17 hours", "Ticketing system — resolved timestamp delta"),
    ("Critical Ticket SLA Breach Rate", "35%", "≤ 10%", "Ticketing system — SLA compliance report"),
    ("At-Risk Customer Identification Time", "7-14 days (manual)", "< 24 hours (automated)", "Risk alert log — timestamp of first alert"),
    ("CSM Manual Health Check Time", "4 hrs/week/CSM", "< 30 min/week/CSM", "Time tracking survey — quarterly"),
    ("MRR Retained from At-Risk Cohort", "Baseline TBD", "+$180K ARR protected", "Revenue retention report"),
    ("Customer NPS (Active Cohort)", "Avg 6.8", "≥ 8.0", "NPS survey — quarterly"),
]
for row_data in metrics:
    table_row(sm_table, row_data)

doc.add_paragraph()

# ── 10. Risk Register ─────────────────────────────────────────────────────────
heading1("10. Risk Register")

risk_table = doc.add_table(rows=1, cols=5)
risk_table.style = "Table Grid"
table_row(risk_table, ["Risk ID", "Risk Description", "Probability", "Impact", "Mitigation Strategy"], header=True)

risks = [
    ("R-01", "AI API rate limits cause latency during peak usage", "Medium", "High", "Implement request queuing + caching layer; fallback to batch processing mode"),
    ("R-02", "Data quality issues in legacy customer records corrupt health scores", "High", "High", "Data validation pipeline at ingestion; null-handling rules defined in BRD appendix"),
    ("R-03", "CS team resistance to adopting new workflow", "Medium", "Medium", "Change management plan; CS team involved in UAT; champions identified early"),
    ("R-04", "Engineering capacity reduced due to competing sprint priorities", "Medium", "High", "Executive sponsor alignment on resource allocation; weekly scope review"),
    ("R-05", "Slack integration approval delayed by IT security review", "Low", "Medium", "Begin IT security review in Week 1; email digest as fallback alert channel"),
    ("R-06", "AI-generated user stories contain inaccurate business context", "Medium", "Low", "All AI outputs reviewed by BA before stakeholder distribution; human-in-loop workflow"),
]
for row_data in risks:
    table_row(risk_table, row_data)

doc.add_paragraph()

# ── 11. Sign-Off ──────────────────────────────────────────────────────────────
heading1("11. Approval & Sign-Off")
body("The undersigned confirm that this BRD accurately represents the business requirements for the BridgeIQ project and approve it for development.")

signoff_table = doc.add_table(rows=1, cols=4)
signoff_table.style = "Table Grid"
table_row(signoff_table, ["Name", "Role", "Signature", "Date"], header=True)

signers = [
    ("_________________", "VP Customer Success (Sponsor)", "_________________", "___________"),
    ("_________________", "Product Manager", "_________________", "___________"),
    ("_________________", "Engineering Lead", "_________________", "___________"),
    ("_________________", "CFO", "_________________", "___________"),
]
for row_data in signers:
    table_row(signoff_table, row_data)

doc.add_paragraph()

# ── Save ─────────────────────────────────────────────────────────────────────
doc.save(OUT)
print(f"BRD saved: {OUT}")
